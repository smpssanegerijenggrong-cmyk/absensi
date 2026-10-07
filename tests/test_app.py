import io
import os
import sys
from datetime import datetime
from zoneinfo import ZoneInfo

# The app is imported only after isolated paths and credentials are configured.
os.environ["SECRET_KEY"]="test-key-not-for-production"
os.environ["ADMIN_USERNAME"]="admin"
os.environ["ADMIN_PASSWORD"]="test-pass"
os.environ["DATABASE_PATH"]="/tmp/sanjara-pytest.sqlite3"
os.environ["UPLOAD_DIR"]="/tmp/sanjara-pytest-surat"
sys.path.insert(0,os.path.abspath(os.path.join(os.path.dirname(__file__),"..")))
import app as mod

def csrf(client):
    with client.session_transaction() as s:
        return s["_csrf"]

def logged_client():
    app=mod.app
    app.config["TESTING"]=True
    c=app.test_client()
    c.get("/login")
    r=c.post("/login",data={"_csrf":csrf(c),"username":"admin","password":"test-pass"})
    assert r.status_code==302
    return c

def new_student(name):
    with mod.connect() as db:
        db.execute("INSERT INTO students(nipd,nisn,name,gender,class_name,qr_token,joined_on) VALUES (?,?,?,?,?,?,?)",
                   ("101","202",name,"L","VII A","qr-"+name,"2026-01-01"))

def configure():
    with mod.connect() as db:
        db.execute("INSERT OR REPLACE INTO config(key,value) VALUES ('latitude','-8.100')")
        db.execute("INSERT OR REPLACE INTO config(key,value) VALUES ('longitude','113.100')")
        db.execute("INSERT OR REPLACE INTO config(key,value) VALUES ('radius_m','200')")
        db.execute("INSERT OR REPLACE INTO config(key,value) VALUES ('workdays','0,1,2,3,4,5,6')")

def test_cutoff():
    tz=ZoneInfo("Asia/Jakarta")
    assert mod.cutoff_status(datetime(2026,10,7,7,0,0,tzinfo=tz),"07:00:00")==("Tepat Waktu",0)
    assert mod.cutoff_status(datetime(2026,10,7,7,0,1,tzinfo=tz),"07:00:00")==("Terlambat",1)
    assert mod.cutoff_status(datetime(2026,10,7,7,12,30,tzinfo=tz),"07:00:00")==("Terlambat",750)

def test_scan_time_duplicate_and_location(monkeypatch):
    c=logged_client()
    configure()
    new_student("StudentForScan")
    tz=ZoneInfo("Asia/Jakarta")
    monkeypatch.setattr(mod,"now_wib",lambda:datetime(2026,10,7,7,12,30,tzinfo=tz))
    payload={"_csrf":csrf(c),"token":"qr-StudentForScan","latitude":"-8.100","longitude":"113.100","accuracy":"10"}
    assert c.post("/scan",data=payload,follow_redirects=True).status_code==200
    with mod.connect() as db:
        st=db.execute("SELECT * FROM students WHERE name='StudentForScan'").fetchone()
        a=db.execute("SELECT * FROM attendance WHERE student_id=? AND day='2026-10-07'",(st["id"],)).fetchone()
        assert a["status"]=="Terlambat"
        assert a["scanned_at"]=="07:12:30"
        assert a["late_seconds"]==750
    c.post("/scan",data=payload)
    with mod.connect() as db:
        assert db.execute("SELECT COUNT(*) FROM attendance WHERE day='2026-10-07' AND source='QR'").fetchone()[0]==1
    new_student("StudentOutside")
    payload.update({"token":"qr-StudentOutside","latitude":"-8.300"})
    c.post("/scan",data=payload)
    with mod.connect() as db:
        assert db.execute("SELECT COUNT(*) FROM attendance WHERE day='2026-10-07' AND source='QR'").fetchone()[0]==1

def test_absence_and_monthly_excel(monkeypatch):
    c=logged_client()
    tz=ZoneInfo("Asia/Jakarta")
    monkeypatch.setattr(mod,"now_wib",lambda:datetime(2026,10,8,16,0,0,tzinfo=tz))
    new_student("StudentForReport")
    with mod.connect() as db:
        s=db.execute("SELECT * FROM students WHERE name='StudentForReport'").fetchone()
        assert mod.row_status(db,s["id"],"2026-10-08")["status"]=="Alpa"
    r=c.get("/report.xlsx?month=2026-10&day=2026-10-08")
    assert r.status_code==200
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(r.data))
    assert wb["REKAP"]["J1"].value=="JUMLAH"
    assert wb["REKAP"]["J2"].value=="SAKIT"
    assert wb["REKAP"]["L2"].value=="ALPA"
    assert "DETAIL WAKTU" in wb.sheetnames

def test_csrf_and_no_default_access():
    c=logged_client()
    assert c.post("/students",data={"name":"Bad"}).status_code==400
    anonymous=mod.app.test_client()
    assert anonymous.get("/students").status_code==302
    assert anonymous.get("/backup.zip").status_code==302
