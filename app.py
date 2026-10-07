import base64
import io
import json
import math
import os
import re
import secrets
import sqlite3
import zipfile
from datetime import date, datetime, time
from functools import wraps
from pathlib import Path
from zoneinfo import ZoneInfo

import qrcode
from flask import Flask, abort, flash, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

TZ = ZoneInfo("Asia/Jakarta")
app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY") or secrets.token_hex(32)
app.config.update(
    MAX_CONTENT_LENGTH=8 * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.getenv("COOKIE_SECURE", "0") == "1",
    DATABASE_PATH=os.getenv("DATABASE_PATH", "data/absensi.sqlite3"),
    UPLOAD_DIR=os.getenv("UPLOAD_DIR", "data/surat"),
)
DB_PATH = Path(app.config["DATABASE_PATH"])
UPLOAD_DIR = Path(app.config["UPLOAD_DIR"])
DB_PATH.parent.mkdir(parents=True, exist_ok=True)
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

def now_wib():
    return datetime.now(TZ)

def connect():
    db = sqlite3.connect(str(DB_PATH), timeout=15)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys=ON")
    db.execute("PRAGMA busy_timeout=5000")
    return db

def init_db():
    with connect() as db:
        db.execute("PRAGMA journal_mode=WAL")
        db.executescript("""
        CREATE TABLE IF NOT EXISTS students (
          id INTEGER PRIMARY KEY, nipd TEXT NOT NULL DEFAULT '', nisn TEXT NOT NULL DEFAULT '',
          name TEXT NOT NULL, gender TEXT NOT NULL CHECK(gender IN ('L','P')),
          class_name TEXT NOT NULL, qr_token TEXT NOT NULL UNIQUE, joined_on TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1);
        CREATE INDEX IF NOT EXISTS idx_students_class ON students(class_name);
        CREATE TABLE IF NOT EXISTS attendance (
          id INTEGER PRIMARY KEY, student_id INTEGER NOT NULL REFERENCES students(id),
          day TEXT NOT NULL, scanned_at TEXT, status TEXT NOT NULL,
          late_seconds INTEGER NOT NULL DEFAULT 0, source TEXT NOT NULL DEFAULT 'QR',
          created_at TEXT NOT NULL, UNIQUE(student_id,day));
        CREATE INDEX IF NOT EXISTS idx_attendance_day ON attendance(day);
        CREATE TABLE IF NOT EXISTS leave_requests (
          id INTEGER PRIMARY KEY, student_id INTEGER NOT NULL REFERENCES students(id),
          day TEXT NOT NULL, kind TEXT NOT NULL CHECK(kind IN ('Ijin','Sakit')),
          reason TEXT NOT NULL, file_path TEXT NOT NULL, state TEXT NOT NULL DEFAULT 'pending',
          created_at TEXT NOT NULL, reviewed_at TEXT);
        CREATE TABLE IF NOT EXISTS holidays(day TEXT PRIMARY KEY, description TEXT NOT NULL DEFAULT '');
        CREATE TABLE IF NOT EXISTS config(key TEXT PRIMARY KEY, value TEXT NOT NULL);
        """)
        for k,v in {'school_name':'SMP SSA Negeri Jenggrong Ranuyoso',
                    'latitude':'','longitude':'','radius_m':'150',
                    'cutoff':'07:00:00','closing':'15:00:00',
                    'workdays':'0,1,2,3,4,5'}.items():
            db.execute("INSERT OR IGNORE INTO config(key,value) VALUES (?,?)",(k,v))

init_db()

def cfg(db):
    return {r["key"]:r["value"] for r in db.execute("SELECT key,value FROM config")}

def is_school_day(db, day, options=None):
    d = date.fromisoformat(day)
    options = options or cfg(db)
    if d.weekday() not in set(int(x) for x in options["workdays"].split(",") if x.strip().isdigit()):
        return False
    return not bool(db.execute("SELECT 1 FROM holidays WHERE day=?", (day,)).fetchone())

def cutoff_status(at, cutoff):
    threshold = datetime.combine(at.date(), time.fromisoformat(cutoff), TZ)
    delta = max(0, int((at-threshold).total_seconds()))
    return ("Terlambat" if delta else "Tepat Waktu", delta)

def distance_m(a,b,c,d):
    p1,p2=math.radians(a),math.radians(c)
    dp=math.radians(c-a)
    dl=math.radians(d-b)
    h=math.sin(dp/2)**2+math.cos(p1)*math.cos(p2)*math.sin(dl/2)**2
    return 6371000*2*math.asin(min(1,math.sqrt(h)))

def login_required(fn):
    @wraps(fn)
    def wrapped(*args,**kwargs):
        if not session.get("admin"):
            return redirect(url_for("login"))
        return fn(*args,**kwargs)
    return wrapped

@app.before_request
def csrf_guard():
    if "_csrf" not in session:
        session["_csrf"]=secrets.token_urlsafe(32)
    if request.method in ("POST","PUT","PATCH","DELETE") and request.endpoint != "static":
        if not secrets.compare_digest(request.form.get("_csrf",""), session["_csrf"]):
            abort(400,"CSRF token invalid")

@app.context_processor
def inject_common():
    return {"csrf":session.get("_csrf",""), "today":now_wib().date().isoformat(),
            "school_name":"SANJARA ABSENSI"}

@app.route("/login",methods=["GET","POST"])
def login():
    if request.method=="POST":
        user=os.getenv("ADMIN_USERNAME","admin")
        password=os.getenv("ADMIN_PASSWORD","")
        if not password or password.startswith("replace-with"):
            flash("Set ADMIN_PASSWORD yang aman di server terlebih dahulu.","error")
        elif secrets.compare_digest(request.form.get("username",""),user) and secrets.compare_digest(request.form.get("password",""),password):
            session.clear()
            session["admin"]=True
            session["_csrf"]=secrets.token_urlsafe(32)
            return redirect(url_for("home"))
        else:
            flash("Nama pengguna atau kata sandi salah.","error")
    return render_template("login.html")

@app.post("/logout")
@login_required
def logout():
    session.clear()
    return redirect(url_for("login"))

def validate_day(value):
    try:
        return date.fromisoformat(value).isoformat()
    except (ValueError,TypeError):
        abort(400,"Tanggal tidak valid")

def month_days(month):
    try:
        y,m=map(int,month.split("-"))
        assert 1<=m<=12 and 2000<=y<=2100
        from calendar import monthrange
        return [date(y,m,d).isoformat() for d in range(1,monthrange(y,m)[1]+1)]
    except (ValueError,AssertionError,AttributeError):
        abort(400,"Bulan tidak valid")

def row_status(db,student_id,day, options=None):
    a=db.execute("SELECT * FROM attendance WHERE student_id=? AND day=?",(student_id,day)).fetchone()
    if a: return dict(a)
    options=options or cfg(db)
    if not is_school_day(db,day,options): return {"status":"Libur","scanned_at":None,"late_seconds":0}
    n=now_wib()
    if day<n.date().isoformat() or (day==n.date().isoformat() and n.time()>=time.fromisoformat(options["closing"])):
        return {"status":"Alpa","scanned_at":None,"late_seconds":0}
    return {"status":"Belum Absen","scanned_at":None,"late_seconds":0}

@app.route("/")
@login_required
def home():
    day=validate_day(request.args.get("day",now_wib().date().isoformat()))
    class_name=request.args.get("class","")
    with connect() as db:
        opts=cfg(db)
        classes=[r[0] for r in db.execute("SELECT DISTINCT class_name FROM students WHERE active=1 ORDER BY class_name")]
        students=db.execute("SELECT * FROM students WHERE active=1 AND (?='' OR class_name=?) ORDER BY class_name,name",(class_name,class_name)).fetchall()
        items=[]
        counters={"Tepat Waktu":0,"Terlambat":0,"Ijin":0,"Sakit":0,"Alpa":0,"Belum Absen":0,"Libur":0}
        for student in students:
            status=row_status(db,student["id"],day,opts)
            counters[status["status"]]=counters.get(status["status"],0)+1
            items.append({"student":student,"attendance":status})
        recent=db.execute("""SELECT a.*,s.name,s.class_name FROM attendance a JOIN students s ON s.id=a.student_id
        WHERE a.day=? ORDER BY a.created_at DESC LIMIT 20""",(day,)).fetchall()
    return render_template("index.html",items=items,counts=counters,day=day,classes=classes,
                           class_name=class_name,options=opts,recent=recent)

@app.route("/students",methods=["GET","POST"])
@login_required
def students():
    if request.method=="POST":
        name=request.form.get("name","").strip()
        gender=request.form.get("gender","").upper()
        class_name=request.form.get("class_name","").strip()
        if not name or not class_name or gender not in ("L","P"):
            flash("Nama, kelas, dan jenis kelamin wajib diisi.","error")
        else:
            with connect() as db:
                db.execute("""INSERT INTO students(nipd,nisn,name,gender,class_name,qr_token,joined_on)
                VALUES (?,?,?,?,?,?,?)""",(request.form.get("nipd","").strip(),request.form.get("nisn","").strip(),
                name,gender,class_name,secrets.token_urlsafe(28),now_wib().date().isoformat()))
            flash("Siswa berhasil ditambahkan.")
        return redirect(url_for("students"))
    with connect() as db:
        rows=db.execute("SELECT * FROM students ORDER BY class_name,name").fetchall()
    return render_template("students.html",students=rows)

@app.post("/students/<int:student_id>/toggle")
@login_required
def toggle_student(student_id):
    with connect() as db:
        cur=db.execute("UPDATE students SET active=1-active WHERE id=?",(student_id,))
        if not cur.rowcount: abort(404)
    return redirect(url_for("students"))

@app.route("/students/import",methods=["POST"])
@login_required
def import_students():
    upload=request.files.get("file")
    if not upload or not upload.filename.lower().endswith(".xlsx"):
        flash("Unggah file .xlsx dengan kolom NIPD, NISN, NAMA, JENIS KELAMIN, KELAS.","error")
        return redirect(url_for("students"))
    try:
        wb=load_workbook(upload.stream,read_only=True,data_only=True)
        sheet=wb.active
        it=sheet.iter_rows(values_only=True)
        header=[str(c or "").strip().upper() for c in next(it)]
        fields=["NIPD","NISN","NAMA","JENIS KELAMIN","KELAS"]
        if any(f not in header for f in fields): raise ValueError("Header tidak sesuai template impor.")
        idx=[header.index(f) for f in fields]
        data=[]
        for row in it:
            vals=[str(row[i] if i<len(row) and row[i] is not None else "").strip() for i in idx]
            if not any(vals):continue
            gender=vals[3].upper()
            gender={"LAKI-LAKI":"L","PEREMPUAN":"P"}.get(gender,gender)
            if gender not in ("L","P") or not vals[2] or not vals[4]:
                raise ValueError("Ada baris dengan nama, kelas, atau jenis kelamin tidak valid.")
            data.append((vals[0],vals[1],vals[2],gender,vals[4],secrets.token_urlsafe(28),now_wib().date().isoformat()))
            if len(data)>5000: raise ValueError("Maksimum 5000 siswa per impor.")
        with connect() as db:
            db.executemany("""INSERT INTO students(nipd,nisn,name,gender,class_name,qr_token,joined_on)
            VALUES (?,?,?,?,?,?,?)""",data)
        flash(f"{len(data)} siswa berhasil diimpor.")
    except Exception as exc:
        flash(f"Impor gagal: {exc}","error")
    return redirect(url_for("students"))

@app.get("/students/template")
@login_required
def import_template():
    wb=Workbook();ws=wb.active;ws.title="IMPORT SISWA"
    ws.append(["NIPD","NISN","NAMA","JENIS KELAMIN","KELAS"])
    ws.append(["","","Contoh Siswa","L","VII A"])
    style_sheet(ws,5)
    bio=io.BytesIO();wb.save(bio);bio.seek(0)
    return send_file(bio,download_name="template_import_sanjara.xlsx",as_attachment=True)

@app.get("/students/<int:student_id>/card")
@login_required
def student_card(student_id):
    with connect() as db:
        student=db.execute("SELECT * FROM students WHERE id=?",(student_id,)).fetchone()
    if not student:abort(404)
    return render_template("card.html",student=student)

@app.get("/students/<int:student_id>/qr.png")
@login_required
def student_qr(student_id):
    with connect() as db:
        student=db.execute("SELECT qr_token FROM students WHERE id=?",(student_id,)).fetchone()
    if not student:abort(404)
    pic=qrcode.make(student["qr_token"])
    bio=io.BytesIO();pic.save(bio,format="PNG");bio.seek(0)
    response=send_file(bio,mimetype="image/png")
    response.headers["Cache-Control"]="no-store"
    return response

@app.get("/scan")
@login_required
def scan():
    return render_template("scan.html")

@app.post("/scan")
@login_required
def record_scan():
    token=request.form.get("token","").strip()
    try:
        lat=float(request.form.get("latitude",""))
        lon=float(request.form.get("longitude",""))
        accuracy=float(request.form.get("accuracy",""))
    except (ValueError,TypeError):
        flash("GPS belum tersedia. Aktifkan izin lokasi perangkat.","error")
        return redirect(url_for("scan"))
    if not (-90<=lat<=90 and -180<=lon<=180) or not math.isfinite(accuracy) or accuracy<0 or accuracy>100:
        flash("Akurasi lokasi tidak mencukupi (maksimum 100 meter).","error")
        return redirect(url_for("scan"))
    with connect() as db:
        opts=cfg(db)
        if not opts["latitude"] or not opts["longitude"]:
            flash("Koordinat sekolah belum diatur.","error")
            return redirect(url_for("scan"))
        if distance_m(lat,lon,float(opts["latitude"]),float(opts["longitude"]))>float(opts["radius_m"]):
            flash("Di luar radius lokasi sekolah. Absensi ditolak.","error")
            return redirect(url_for("scan"))
        student=db.execute("SELECT * FROM students WHERE qr_token=? AND active=1",(token,)).fetchone()
        if not student:
            flash("QR tidak ditemukan atau siswa tidak aktif.","error")
            return redirect(url_for("scan"))
        at=now_wib()
        day=at.date().isoformat()
        if not is_school_day(db,day,opts):
            flash("Hari ini bukan hari sekolah.","error")
            return redirect(url_for("scan"))
        existing=db.execute("SELECT id FROM attendance WHERE student_id=? AND day=?",(student["id"],day)).fetchone()
        if existing:
            flash(f"{student['name']} sudah tercatat absen hari ini.","error")
            return redirect(url_for("scan"))
        status,late=cutoff_status(at,opts["cutoff"])
        try:
            db.execute("""INSERT INTO attendance(student_id,day,scanned_at,status,late_seconds,source,created_at)
            VALUES (?,?,?,?,?,'QR',?)""",(student["id"],day,at.strftime("%H:%M:%S"),status,late,at.isoformat()))
        except sqlite3.IntegrityError:
            flash("Siswa sudah tercatat hari ini.","error")
            return redirect(url_for("scan"))
    flash(f"{student['name']} • {student['class_name']} • {at.strftime('%H:%M:%S')} WIB • {status} ({late//60} menit {late%60} detik).")
    return redirect(url_for("scan"))

@app.route("/leave",methods=["GET","POST"])
@login_required
def leave():
    if request.method=="POST":
        try:
            student_id=int(request.form.get("student_id",""))
            day=validate_day(request.form.get("day",""))
            kind=request.form.get("kind","")
            reason=request.form.get("reason","").strip()
            file=request.files.get("letter")
            if kind not in ("Ijin","Sakit") or not reason or not file or not file.filename:
                raise ValueError("Alasan dan surat izin wajib diisi.")
            extension=Path(secure_filename(file.filename)).suffix.lower()
            if extension not in (".pdf",".jpg",".jpeg",".png"):
                raise ValueError("Hanya PDF/JPG/PNG yang diterima.")
            content=file.read()
            # Inspect signatures, not merely extension
            valid=(extension==".pdf" and content.startswith(b"%PDF-")) or (
                extension in (".jpg",".jpeg") and content.startswith(b"\xff\xd8\xff")) or (
                extension==".png" and content.startswith(b"\x89PNG\r\n\x1a\n"))
            if not valid: raise ValueError("Format isi berkas tidak sesuai ekstensi.")
            with connect() as db:
                if not db.execute("SELECT 1 FROM students WHERE id=? AND active=1",(student_id,)).fetchone():
                    raise ValueError("Siswa tidak ditemukan.")
                if db.execute("SELECT 1 FROM attendance WHERE student_id=? AND day=?",(student_id,day)).fetchone():
                    raise ValueError("Absensi tanggal ini sudah tercatat.")
                filename=secrets.token_hex(20)+extension
                path=UPLOAD_DIR/filename
                path.write_bytes(content)
                try:
                    db.execute("""INSERT INTO leave_requests(student_id,day,kind,reason,file_path,state,created_at)
                    VALUES (?,?,?,?,?,'pending',?)""",(student_id,day,kind,reason,filename,now_wib().isoformat()))
                except Exception:
                    path.unlink(missing_ok=True)
                    raise
            flash("Pengajuan tersimpan, menunggu persetujuan.")
        except (ValueError,sqlite3.Error) as exc:
            flash(str(exc),"error")
        return redirect(url_for("leave"))
    with connect() as db:
        students=db.execute("SELECT * FROM students WHERE active=1 ORDER BY class_name,name").fetchall()
        rows=db.execute("""SELECT l.*,s.name,s.class_name FROM leave_requests l
        JOIN students s ON s.id=l.student_id ORDER BY l.created_at DESC LIMIT 200""").fetchall()
    return render_template("leave.html",students=students,requests=rows)

@app.post("/leave/<int:request_id>/review")
@login_required
def review_leave(request_id):
    action=request.form.get("action")
    if action not in ("approve","reject"):abort(400)
    with connect() as db:
        item=db.execute("SELECT * FROM leave_requests WHERE id=?",(request_id,)).fetchone()
        if not item:abort(404)
        if item["state"]!="pending":
            flash("Pengajuan sudah diproses.","error")
        elif action=="reject":
            db.execute("UPDATE leave_requests SET state='rejected',reviewed_at=? WHERE id=?",(now_wib().isoformat(),request_id))
            flash("Pengajuan ditolak.")
        else:
            try:
                db.execute("""INSERT INTO attendance(student_id,day,scanned_at,status,late_seconds,source,created_at)
                VALUES (?,?,NULL,?,0,'SURAT',?)""",(item["student_id"],item["day"],item["kind"],now_wib().isoformat()))
                db.execute("UPDATE leave_requests SET state='approved',reviewed_at=? WHERE id=?",(now_wib().isoformat(),request_id))
                flash("Pengajuan disetujui dan rekap diperbarui.")
            except sqlite3.IntegrityError:
                flash("Tanggal ini sudah memiliki absensi. Pengajuan tidak diproses.","error")
    return redirect(url_for("leave"))

@app.get("/leave/<int:request_id>/letter")
@login_required
def letter(request_id):
    with connect() as db:
        row=db.execute("SELECT file_path FROM leave_requests WHERE id=?",(request_id,)).fetchone()
    if not row:abort(404)
    path=UPLOAD_DIR/Path(row["file_path"]).name
    if not path.exists():abort(404)
    return send_file(path,as_attachment=True,download_name="surat_izin"+path.suffix)

@app.route("/settings",methods=["GET","POST"])
@login_required
def settings():
    if request.method=="POST":
        try:
            lat=float(request.form.get("latitude",""))
            lon=float(request.form.get("longitude",""))
            radius=int(request.form.get("radius_m",""))
            cutoff=time.fromisoformat(request.form.get("cutoff",""))
            closing=time.fromisoformat(request.form.get("closing",""))
            days=request.form.getlist("workdays")
            if not (-90<=lat<=90 and -180<=lon<=180 and 20<=radius<=2000):
                raise ValueError("Koordinat/radius tidak valid (20–2000 meter).")
            if closing<=cutoff or not days or any(x not in list(map(str,range(7))) for x in days):
                raise ValueError("Jam akhir harus setelah batas hadir, dan hari sekolah harus dipilih.")
            values={"latitude":str(lat),"longitude":str(lon),"radius_m":str(radius),
                    "cutoff":cutoff.strftime("%H:%M:%S"),"closing":closing.strftime("%H:%M:%S"),
                    "workdays":",".join(sorted(set(days)))}
            with connect() as db:
                for k,v in values.items():db.execute("INSERT OR REPLACE INTO config(key,value) VALUES (?,?)",(k,v))
            flash("Pengaturan disimpan.")
        except (ValueError,TypeError) as exc:
            flash(str(exc),"error")
        return redirect(url_for("settings"))
    with connect() as db:
        options=cfg(db)
        holidays=db.execute("SELECT * FROM holidays ORDER BY day DESC LIMIT 100").fetchall()
    return render_template("settings.html",options=options,holidays=holidays)

@app.post("/settings/holiday")
@login_required
def add_holiday():
    day=validate_day(request.form.get("day",""))
    desc=request.form.get("description","").strip()
    with connect() as db:
        db.execute("INSERT OR REPLACE INTO holidays(day,description) VALUES (?,?)",(day,desc))
    flash("Hari libur disimpan.")
    return redirect(url_for("settings"))

@app.post("/settings/holiday/<day>/delete")
@login_required
def delete_holiday(day):
    day=validate_day(day)
    with connect() as db:db.execute("DELETE FROM holidays WHERE day=?",(day,))
    return redirect(url_for("settings"))

def report_rows(db, month, selected_day, class_name):
    days=month_days(month)
    options=cfg(db)
    students=db.execute("""SELECT * FROM students WHERE active=1 AND joined_on<=?
      AND (?='' OR class_name=?) ORDER BY class_name,name""",(days[-1],class_name,class_name)).fetchall()
    output=[]
    details=[]
    for student in students:
        counts={"Sakit":0,"Ijin":0,"Alpa":0}
        for day in days:
            if day<student["joined_on"] or day>now_wib().date().isoformat():continue
            state=row_status(db,student["id"],day,options)
            if state["status"] in counts:counts[state["status"]]+=1
            if state["status"] in ("Tepat Waktu","Terlambat","Sakit","Ijin"):
                details.append([day,student["nipd"],student["nisn"],student["name"],student["class_name"],
                    state.get("scanned_at") or "",state["status"],
                    (state.get("late_seconds") or 0)//60, state.get("source") or ""])
        current=row_status(db,student["id"],selected_day,options) if selected_day>=student["joined_on"] else {"status":"Belum Terdaftar"}
        output.append((student,current,counts))
    return output, details

@app.get("/report")
@login_required
def report():
    day=validate_day(request.args.get("day",now_wib().date().isoformat()))
    month=request.args.get("month",day[:7])
    month_days(month)
    if not day.startswith(month):
        day=month+"-01"
    class_name=request.args.get("class","")
    with connect() as db:
        records,details=report_rows(db,month,day,class_name)
        classes=[r[0] for r in db.execute("SELECT DISTINCT class_name FROM students WHERE active=1 ORDER BY class_name")]
    return render_template("report.html",records=records,day=day,month=month,class_name=class_name,classes=classes)

def style_sheet(ws,columns):
    ws.freeze_panes="A3"
    for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,2)):
        for cell in row:
            cell.fill=PatternFill("solid",fgColor="134E4A")
            cell.font=Font(bold=True,color="FFFFFF",size=10)
            cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    widths=[7,16,20,30,20,14]+[15]*max(0,columns-6)
    from openpyxl.utils import get_column_letter
    for i in range(1,columns+1):ws.column_dimensions[get_column_letter(i)].width=widths[i-1]
    ws.row_dimensions[1].height=27
    ws.row_dimensions[2].height=27

@app.get("/report.xlsx")
@login_required
def export_report():
    day=validate_day(request.args.get("day",now_wib().date().isoformat()))
    month=request.args.get("month",day[:7])
    month_days(month)
    if not day.startswith(month):day=month+"-01"
    class_name=request.args.get("class","")
    with connect() as db:
        records,details=report_rows(db,month,day,class_name)
        holidays=db.execute("SELECT * FROM holidays WHERE day BETWEEN ? AND ? ORDER BY day",(month+"-01",month+"-31")).fetchall()
    wb=Workbook()
    ws=wb.active;ws.title="REKAP"
    ws.append(["No","NIPD","NISN ","NAMA","JENIS KELAMIN","KELAS","SAKIT","IJIN","ALPA","JUMLAH",None,None])
    ws.merge_cells("J1:L1")
    ws.append([None,None,None,None,None,None,None,None,None,"SAKIT","IJIN","ALPA"])
    for i,(s,state,c) in enumerate(records,1):
        ws.append([i,s["nipd"],s["nisn"],s["name"],s["gender"],s["class_name"],
            int(state["status"]=="Sakit"),int(state["status"]=="Ijin"),int(state["status"]=="Alpa"),
            c["Sakit"],c["Ijin"],c["Alpa"]])
    style_sheet(ws,12)
    detail=wb.create_sheet("DETAIL WAKTU")
    detail.append(["Tanggal","NIPD","NISN","Nama","Kelas","Jam Scan WIB","Status","Menit Terlambat","Sumber"])
    for row in details:detail.append(row)
    detail.freeze_panes="A2"
    for c in detail[1]:
        c.fill=PatternFill("solid",fgColor="134E4A");c.font=Font(color="FFFFFF",bold=True)
    for col,width in zip("ABCDEFGHI",[17,18,19,30,16,19,18,20,16]):detail.column_dimensions[col].width=width
    legend=wb.create_sheet("LEGENDA")
    for row in [["SANJARA ABSENSI — PETUNJUK REKAP"],["Tanggal dipilih",day],["Bulan rekap",month],
                ["Kolom G–I","Status SAKIT, IJIN, ALPA pada tanggal yang dipilih (1 atau 0)"],
                ["Kolom J–L","Total SAKIT, IJIN, ALPA selama bulan rekap"],
                ["Terlambat","Tetap hadir, waktu scan tersimpan pada sheet DETAIL WAKTU"],
                ["Alpa","Dihitung hanya setelah penutupan hari sekolah, tanpa absensi/izin disetujui"],
                ["Waktu","WIB / Asia/Jakarta"]]:legend.append(row)
    legend.column_dimensions["A"].width=30;legend.column_dimensions["B"].width=86
    holiday=wb.create_sheet("KALENDER")
    holiday.append(["Hari Libur","Keterangan"])
    for h in holidays:holiday.append([h["day"],h["description"]])
    holiday.column_dimensions["A"].width=20;holiday.column_dimensions["B"].width=46
    bio=io.BytesIO();wb.save(bio);bio.seek(0)
    return send_file(bio,as_attachment=True,download_name=f"rekap_sanjara_{month}_{day}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

def build_backup():
    # SQLite online backup avoids copying an inconsistent WAL database.
    temp=io.BytesIO()
    import tempfile
    with tempfile.TemporaryDirectory() as directory:
        db_snapshot=Path(directory)/"absensi.sqlite3"
        with connect() as source, sqlite3.connect(db_snapshot) as target:
            source.backup(target)
        with zipfile.ZipFile(temp,"w",zipfile.ZIP_DEFLATED) as z:
            z.write(db_snapshot,"absensi.sqlite3")
            for f in UPLOAD_DIR.iterdir():
                if f.is_file() and not f.is_symlink():
                    z.write(f,"surat/"+f.name)
    temp.seek(0)
    return temp

@app.get("/backup.zip")
@login_required
def download_backup():
    return send_file(build_backup(),as_attachment=True,
                     download_name="backup_sanjara_"+now_wib().strftime("%Y%m%d_%H%M")+".zip")

@app.post("/backup/drive")
@login_required
def backup_drive():
    service=os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON","")
    folder=os.getenv("DRIVE_FOLDER_ID","")
    if not service or not folder:
        flash("Konfigurasi GOOGLE_SERVICE_ACCOUNT_JSON dan DRIVE_FOLDER_ID belum tersedia.","error")
        return redirect(url_for("settings"))
    try:
        from google.oauth2 import service_account
        from google.auth.transport.requests import AuthorizedSession
        credentials=service_account.Credentials.from_service_account_info(
            json.loads(service),scopes=["https://www.googleapis.com/auth/drive.file"])
        session_drive=AuthorizedSession(credentials)
        import requests
        metadata={"name":"backup_sanjara_"+now_wib().strftime("%Y%m%d_%H%M%S")+".zip","parents":[folder]}
        boundary="sanjara"+secrets.token_hex(16)
        zipdata=build_backup().getvalue()
        body=(("--"+boundary+"\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n").encode()+
              json.dumps(metadata).encode()+("\r\n--"+boundary+"\r\nContent-Type: application/zip\r\n\r\n").encode()+
              zipdata+("\r\n--"+boundary+"--\r\n").encode())
        resp=session_drive.post("https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart",
            data=body,headers={"Content-Type":"multipart/related; boundary="+boundary},timeout=60)
        resp.raise_for_status()
        flash("Backup berhasil diunggah ke Google Drive.")
    except Exception as exc:
        app.logger.exception("Drive backup failed")
        flash("Backup Drive gagal. Periksa akses folder, kredensial, dan koneksi server.","error")
    return redirect(url_for("settings"))

@app.get("/health")
def health():
    return {"status":"ok","app":"SANJARA ABSENSI"}

if __name__=="__main__":
    app.run(host="0.0.0.0",port=int(os.getenv("PORT",5000)),debug=False)
